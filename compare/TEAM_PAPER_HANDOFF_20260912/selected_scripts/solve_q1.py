#!/usr/bin/env python3
"""Q1 baseline: detect time-frequency conflicts in D题 plans.

The script reads the unchanged raw workbook, expands repeated uses, applies
half-open interval intersection, independently recomputes conflicts through
discrete occupied cells, and writes reviewable Q1 outputs.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import json
from pathlib import Path
import re
from typing import Iterable

from openpyxl import load_workbook


INTERVAL_RE = re.compile(r"\s*\[\s*(-?\d+)\s*,\s*(-?\d+)\s*\)\s*")
EXPECTED_HEADERS = ["用频装备编号", "频段区间", "时间区间", "间隔时长", "使用次数"]


def parse_interval(value: object, field: str, row: int) -> tuple[int, int]:
    if not isinstance(value, str):
        raise ValueError(f"{field} row {row}: interval is not text: {value!r}")
    match = INTERVAL_RE.fullmatch(value)
    if not match:
        raise ValueError(f"{field} row {row}: invalid half-open interval: {value!r}")
    left, right = map(int, match.groups())
    if right <= left:
        raise ValueError(f"{field} row {row}: non-positive interval: {value!r}")
    return left, right


def read_plans(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"unexpected sheets in {path}: {workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    headers = [cell.value for cell in sheet[1]]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"unexpected headers: {headers!r}")

    plans: list[dict] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or value == "" for value in values):
            continue
        record = dict(zip(headers, values))
        identifier = record["用频装备编号"]
        if not isinstance(identifier, str) or not re.fullmatch(r"[ABC]\d{3}", identifier):
            raise ValueError(f"row {row_number}: invalid equipment id: {identifier!r}")
        frequency = parse_interval(record["频段区间"], "频段区间", row_number)
        time = parse_interval(record["时间区间"], "时间区间", row_number)
        gap = record["间隔时长"]
        uses = record["使用次数"]
        if not isinstance(gap, int) or gap <= 0:
            raise ValueError(f"row {row_number}: invalid interval length: {gap!r}")
        if not isinstance(uses, int) or uses <= 0:
            raise ValueError(f"row {row_number}: invalid usage count: {uses!r}")
        if frequency[0] < 0 or frequency[1] > 100:
            raise ValueError(f"row {row_number}: frequency interval outside [0,100): {frequency}")
        plans.append(
            {
                "id": identifier,
                "category": identifier[0],
                "frequency": frequency,
                "time": time,
                "gap": gap,
                "uses": uses,
            }
        )
    if len(plans) != 150:
        raise ValueError(f"expected 150 plans, found {len(plans)}")
    if len({plan["id"] for plan in plans}) != len(plans):
        raise ValueError("equipment ids are not unique")
    return plans


def expanded_events(plan: dict) -> list[tuple[int, int, int, int]]:
    f0, f1 = plan["frequency"]
    t0, t1 = plan["time"]
    duration = t1 - t0
    start_step = duration + plan["gap"]
    return [(f0, f1, t0 + k * start_step, t1 + k * start_step) for k in range(plan["uses"])]


def validate_recurrence_examples(plans: list[dict]) -> None:
    by_id = {plan["id"]: plan for plan in plans}
    a001_events = expanded_events(by_id["A001"])
    if a001_events[1] != (80, 90, 100, 105):
        raise AssertionError(
            "A001 recurrence disagrees with PDF page 2: expected second event [100,105)"
        )
    c083_events = expanded_events(by_id["C083"])
    if c083_events[-1] != (71, 74, 641, 643):
        raise AssertionError(
            "C083 recurrence boundary mismatch: expected twelfth event [641,643)"
        )


def overlaps(left: tuple[int, int], right: tuple[int, int]) -> bool:
    return max(left[0], right[0]) < min(left[1], right[1])


def interval_conflicts(plans: list[dict]) -> tuple[set[tuple[str, str]], Counter[tuple[str, str]]]:
    conflicts: set[tuple[str, str]] = set()
    event_hits: Counter[tuple[str, str]] = Counter()
    events = {plan["id"]: expanded_events(plan) for plan in plans}
    frequencies = {plan["id"]: plan["frequency"] for plan in plans}
    for index, left in enumerate(plans):
        for right in plans[index + 1 :]:
            if not overlaps(frequencies[left["id"]], frequencies[right["id"]]):
                continue
            for lf0, lf1, lt0, lt1 in events[left["id"]]:
                for rf0, rf1, rt0, rt1 in events[right["id"]]:
                    if overlaps((lt0, lt1), (rt0, rt1)):
                        pair = tuple(sorted((left["id"], right["id"])))
                        conflicts.add(pair)
                        event_hits[pair] += 1
    return conflicts, event_hits


def occupied_cells(plan: dict) -> set[tuple[int, int]]:
    f0, f1 = plan["frequency"]
    t0, t1 = plan["time"]
    duration = t1 - t0
    start_step = duration + plan["gap"]
    cells: set[tuple[int, int]] = set()
    for k in range(plan["uses"]):
        shifted_start = t0 + k * start_step
        shifted_end = t1 + k * start_step
        cells.update((time, frequency) for time in range(shifted_start, shifted_end) for frequency in range(f0, f1))
    return cells


def cell_conflicts(plans: list[dict]) -> set[tuple[str, str]]:
    occupied = {plan["id"]: occupied_cells(plan) for plan in plans}
    conflicts: set[tuple[str, str]] = set()
    for index, left in enumerate(plans):
        for right in plans[index + 1 :]:
            if occupied[left["id"]].isdisjoint(occupied[right["id"]]):
                continue
            conflicts.add(tuple(sorted((left["id"], right["id"]))))
    return conflicts


def pair_category(pair: tuple[str, str]) -> str:
    return "-".join(sorted((pair[0][0], pair[1][0])))


def connected_components(plan_ids: list[str], conflicts: set[tuple[str, str]]) -> list[list[str]]:
    adjacency: dict[str, set[str]] = {identifier: set() for identifier in plan_ids}
    for left, right in conflicts:
        adjacency[left].add(right)
        adjacency[right].add(left)
    components: list[list[str]] = []
    unseen = set(plan_ids)
    while unseen:
        root = min(unseen)
        stack = [root]
        component: list[str] = []
        unseen.remove(root)
        while stack:
            current = stack.pop()
            component.append(current)
            neighbours = adjacency[current] & unseen
            unseen.difference_update(neighbours)
            stack.extend(sorted(neighbours, reverse=True))
        components.append(sorted(component))
    return sorted(components, key=lambda component: (-len(component), component[0]))


def write_outputs(plans: list[dict], conflicts: set[tuple[str, str]], event_hits: Counter[tuple[str, str]], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    plot_data_dir = output_dir / "plot_data"
    plot_data_dir.mkdir(parents=True, exist_ok=True)
    ordered_pairs = sorted(conflicts)
    involved = {identifier for pair in conflicts for identifier in pair}
    category_counts = Counter(pair_category(pair) for pair in ordered_pairs)
    degree = Counter(identifier for pair in ordered_pairs for identifier in pair)
    plan_ids = [plan["id"] for plan in plans]
    components = connected_components(plan_ids, conflicts)
    all_events = [event for plan in plans for event in expanded_events(plan)]
    category_sizes = Counter(plan["category"] for plan in plans)
    category_pairs = ("A-A", "A-B", "A-C", "B-B", "B-C", "C-C")
    possible_by_category = {}
    rate_by_category = {}
    for category_pair in category_pairs:
        left_category, right_category = category_pair.split("-")
        if left_category == right_category:
            possible = category_sizes[left_category] * (category_sizes[left_category] - 1) // 2
        else:
            possible = category_sizes[left_category] * category_sizes[right_category]
        possible_by_category[category_pair] = possible
        rate_by_category[category_pair] = category_counts.get(category_pair, 0) / possible if possible else 0
    event_overlap_distribution = Counter(event_hits.values())
    per_category = {}
    for category in "ABC":
        total = sum(plan["category"] == category for plan in plans)
        used = sum(identifier[0] == category for identifier in involved)
        per_category[category] = {"total_plans": total, "involved_plans": used, "uninvolved_plans": total - used}

    result = {
        "question": "D-Q1",
        "method": "plan expansion + half-open interval intersection",
        "assumptions": [
            "If duration=t1-t0 and idle_gap is the interval after one use ends, usage k occupies [t0+k*(duration+idle_gap), t1+k*(duration+idle_gap)).",
            "Two plans conflict when at least one pair of usage events overlaps in both frequency and time.",
            "A plan pair is listed once even if multiple usage-event pairs overlap.",
        ],
        "plan_count": len(plans),
        "expanded_event_count": len(all_events),
        "candidate_plan_pair_count": len(plans) * (len(plans) - 1) // 2,
        "conflict_pair_count": len(ordered_pairs),
        "involved_plan_count": len(involved),
        "conflict_pairs_by_category": dict(sorted(category_counts.items())),
        "possible_pairs_by_category": possible_by_category,
        "conflict_rate_by_category": rate_by_category,
        "plans_by_category": per_category,
        "event_overlap_count": sum(event_hits.values()),
        "event_overlap_count_distribution": dict(sorted(event_overlap_distribution.items())),
        "max_event_overlap_count_per_pair": max(event_hits.values(), default=0),
        "max_conflict_degree": max(degree.values(), default=0),
        "max_conflict_degree_plan_ids": sorted(identifier for identifier, value in degree.items() if value == max(degree.values(), default=0)),
        "mean_conflict_degree": (2 * len(ordered_pairs) / len(plans)) if plans else 0,
        "time_envelope": [min(event[2] for event in all_events), max(event[3] for event in all_events)],
        "frequency_envelope": [min(event[0] for event in all_events), max(event[1] for event in all_events)],
        "connected_component_count": len(components),
        "largest_connected_component_size": len(components[0]) if components else 0,
        "isolated_plan_count": sum(len(component) == 1 for component in components),
        "isolated_plan_ids": sorted(component[0] for component in components if len(component) == 1),
        "conflict_pairs": [{"序号": i, "冲突装备1": left, "冲突装备2": right, "event_overlap_count": event_hits[(left, right)]} for i, (left, right) in enumerate(ordered_pairs, start=1)],
        "validation": {
            "pdf_A001_recurrence_example": "PASS",
            "C083_last_event_boundary": "PASS",
            "independent_cell_recomputation": "PASS",
            "pair_order_unique": len(ordered_pairs) == len(set(ordered_pairs)),
            "all_ids_in_input": involved <= {plan["id"] for plan in plans},
        },
    }
    (output_dir / "results.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    with (output_dir / "conflict_pairs.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["序号", "冲突装备1", "冲突装备2", "事件重叠次数"])
        for row in result["conflict_pairs"]:
            writer.writerow([row["序号"], row["冲突装备1"], row["冲突装备2"], row["event_overlap_count"]])
    with (plot_data_dir / "conflict_counts_by_category.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["类别组合", "可配对数", "冲突对数", "冲突率"])
        for category_pair in category_pairs:
            writer.writerow([
                category_pair,
                possible_by_category[category_pair],
                category_counts.get(category_pair, 0),
                rate_by_category[category_pair],
            ])
    with (plot_data_dir / "degree_by_plan.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["装备编号", "类别", "冲突度", "是否涉及冲突"])
        for plan in plans:
            identifier = plan["id"]
            writer.writerow([identifier, plan["category"], degree[identifier], int(identifier in involved)])
    with (plot_data_dir / "component_sizes.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["连通分量序号", "装备数", "装备编号"])
        for index, component in enumerate(components, start=1):
            writer.writerow([index, len(component), ",".join(component)])
    with (plot_data_dir / "event_overlap_distribution.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["单个冲突对的事件重叠次数", "冲突对数"])
        for overlap_count, pair_count in sorted(event_overlap_distribution.items()):
            writer.writerow([overlap_count, pair_count])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data/raw/D题/附件/附件1.xlsx"))
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/q1"))
    args = parser.parse_args()

    plans = read_plans(args.input)
    validate_recurrence_examples(plans)
    interval_pairs, event_hits = interval_conflicts(plans)
    cell_pairs = cell_conflicts(plans)
    if interval_pairs != cell_pairs:
        only_interval = sorted(interval_pairs - cell_pairs)[:10]
        only_cells = sorted(cell_pairs - interval_pairs)[:10]
        raise AssertionError(f"independent conflict recomputation mismatch; interval-only={only_interval}, cell-only={only_cells}")
    write_outputs(plans, interval_pairs, event_hits, args.output_dir)
    print(json.dumps({
        "status": "PASS",
        "plans": len(plans),
        "conflict_pairs": len(interval_pairs),
        "event_overlap_count": sum(event_hits.values()),
        "output_dir": str(args.output_dir),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
