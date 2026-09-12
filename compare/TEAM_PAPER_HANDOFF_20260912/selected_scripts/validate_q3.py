#!/usr/bin/env python3
"""Independently validate the fixed-Q2 D-Q3 schedule and result workbook."""
from __future__ import annotations

import argparse
from collections import defaultdict
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from scripts.solve_q1 import read_plans


INTERVAL_RE = re.compile(r"\s*\[\s*(-?\d+)\s*,\s*(-?\d+)\s*\)\s*")
C_WIDTH = 3
C_DURATION = 2
C_GAP = 8
C_USES = 12
C_STEP = C_DURATION + C_GAP


def parse_interval(value: object, field: str) -> tuple[int, int]:
    match = INTERVAL_RE.fullmatch(str(value or ""))
    if not match:
        raise ValueError(f"invalid {field} interval: {value!r}")
    left, right = map(int, match.groups())
    if right <= left:
        raise ValueError(f"non-positive {field} interval: {value!r}")
    return left, right


def overlaps(left: tuple[int, int], right: tuple[int, int]) -> bool:
    return max(left[0], right[0]) < min(left[1], right[1])


def expand_plan(plan: dict) -> list[tuple[int, int, int, int]]:
    f0, f1 = plan["frequency"]
    t0, t1 = plan["time"]
    duration = t1 - t0
    step = duration + plan["gap"]
    return [
        (f0, f1, t0 + k * step, t1 + k * step)
        for k in range(plan["uses"])
    ]


def cells(events: list[tuple[int, int, int, int]]) -> set[tuple[int, int]]:
    return {
        (time, frequency)
        for f0, f1, t0, t1 in events
        for time in range(t0, t1)
        for frequency in range(f0, f1)
    }


def apply_q2_actions(plans: list[dict], actions: list[dict]) -> tuple[list[dict], list[str]]:
    by_id = {plan["id"]: dict(plan) for plan in plans}
    seen: set[str] = set()
    errors: list[str] = []
    adjusted: list[dict] = []
    for row in actions:
        identifier = row.get("装备编号")
        if identifier in seen:
            errors.append(f"duplicate Q2 action: {identifier}")
            continue
        seen.add(identifier)
        if identifier not in by_id:
            errors.append(f"unknown Q2 equipment id: {identifier}")
            continue
        raw = by_id[identifier]
        action = row.get("动作")
        try:
            if action == "revoke":
                adjusted.append({**raw, "revoked": True})
            elif action == "freq":
                df = row.get("频移")
                dt = row.get("时移")
                frequency = parse_interval(row.get("频段区间"), "Q2 frequency")
                if not isinstance(df, int) or df == 0 or dt != 0:
                    raise ValueError("invalid frequency shift fields")
                if frequency != (raw["frequency"][0] + df, raw["frequency"][1] + df):
                    raise ValueError("frequency interval does not match shift")
                adjusted.append({**raw, "frequency": frequency, "revoked": False})
            elif action == "time":
                df = row.get("频移")
                dt = row.get("时移")
                time = parse_interval(row.get("时间区间"), "Q2 time")
                if not isinstance(dt, int) or dt == 0 or df != 0:
                    raise ValueError("invalid time shift fields")
                if time != (raw["time"][0] + dt, raw["time"][1] + dt):
                    raise ValueError("time interval does not match shift")
                adjusted.append({**raw, "time": time, "revoked": False})
            else:
                raise ValueError(f"unsupported action {action!r}")
        except (TypeError, ValueError) as exc:
            errors.append(f"Q2 action {identifier}: {exc}")
    for identifier, raw in by_id.items():
        if identifier not in seen:
            adjusted.append({**raw, "revoked": False})
    return adjusted, errors


def pair_conflicts(named_plans: list[tuple[str, dict]]) -> list[tuple[str, str]]:
    events = [(identifier, plan, expand_plan(plan)) for identifier, plan in named_plans]
    conflicts: list[tuple[str, str]] = []
    for left_index, (left_id, _left_plan, left_events) in enumerate(events):
        for right_id, _right_plan, right_events in events[left_index + 1 :]:
            if any(
                overlaps((left_event[0], left_event[1]), (right_event[0], right_event[1]))
                and overlaps((left_event[2], left_event[3]), (right_event[2], right_event[3]))
                for left_event in left_events
                for right_event in right_events
            ):
                conflicts.append(tuple(sorted((left_id, right_id))))
    return sorted(set(conflicts))


def check_workbook(workbook_path: Path, expected_rows: list[list[object]]) -> dict:
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        sheets = workbook.sheetnames
        if sheets != ["Sheet1"]:
            return {"passed": False, "errors": [f"unexpected sheets: {sheets}"]}
        sheet = workbook["Sheet1"]
        headers = [sheet.cell(1, index).value for index in range(1, 4)]
        rows = [
            list(row)
            for row in sheet.iter_rows(min_row=2, max_col=3, values_only=True)
        ]
        nonempty = [row for row in rows if any(value is not None for value in row)]
        errors = []
        if headers != ["新增用频装备序号", "调整后频段区间", "调整后时间区间"]:
            errors.append(f"unexpected headers: {headers}")
        if nonempty != expected_rows:
            errors.append("workbook rows do not exactly match Q3 results.json")
        return {
            "passed": not errors,
            "errors": errors,
            "sheet_names": sheets,
            "headers": headers,
            "nonempty_row_count": len(nonempty),
            "max_row": sheet.max_row,
        }
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data/raw/D题/附件/附件1.xlsx"))
    parser.add_argument("--q2-result", type=Path, default=Path("outputs/q2/results.json"))
    parser.add_argument("--q3-result", type=Path, default=Path("outputs/q3/results.json"))
    parser.add_argument("--workbook", type=Path, default=Path("outputs/q3/result3.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("outputs/q3/validation.json"))
    parser.add_argument("--horizon", type=int, default=643)
    args = parser.parse_args()

    errors: list[str] = []
    raw_plans = read_plans(args.input)
    q2_result = json.loads(args.q2_result.read_text(encoding="utf-8"))
    q3_result = json.loads(args.q3_result.read_text(encoding="utf-8"))

    q2_plans, q2_action_errors = apply_q2_actions(
        raw_plans, q2_result.get("summary", {}).get("actions", [])
    )
    errors.extend(q2_action_errors)
    q2_active = [plan for plan in q2_plans if not plan.get("revoked", False)]
    q2_named = [(plan["id"], plan) for plan in q2_active]

    q2_boundary_errors = []
    q2_cells: set[tuple[int, int]] = set()
    for identifier, plan in q2_named:
        for event in expand_plan(plan):
            f0, f1, t0, t1 = event
            if f0 < 0 or f1 > 100 or t0 < 0 or t1 > args.horizon:
                q2_boundary_errors.append({"id": identifier, "event": event})
        q2_cells.update(cells(expand_plan(plan)))
    q2_conflicts = pair_conflicts(q2_named)

    selected = q3_result.get("selected_plans", [])
    expected_rows = [
        [row.get("新增用频装备序号"), row.get("频段区间"), row.get("时间区间")]
        for row in selected
    ]
    new_named: list[tuple[str, dict]] = []
    sequence_errors = []
    expected_sequences = list(range(1, len(selected) + 1))
    actual_sequences = [row.get("新增用频装备序号") for row in selected]
    if actual_sequences != expected_sequences:
        sequence_errors.append("new plan sequence is not contiguous from 1")

    new_boundary_errors = []
    template_errors = []
    new_cells: set[tuple[int, int]] = set()
    for row in selected:
        sequence = row.get("新增用频装备序号")
        identifier = f"N{int(sequence):03d}" if isinstance(sequence, int) else f"N?{sequence}"
        try:
            frequency = parse_interval(row.get("频段区间"), "Q3 frequency")
            time = parse_interval(row.get("时间区间"), "Q3 time")
            if frequency[1] - frequency[0] != C_WIDTH:
                template_errors.append({"sequence": sequence, "field": "frequency_width"})
            if time[1] - time[0] != C_DURATION:
                template_errors.append({"sequence": sequence, "field": "duration"})
            plan = {
                "frequency": frequency,
                "time": time,
                "gap": C_GAP,
                "uses": C_USES,
            }
            events = expand_plan(plan)
            for event in events:
                f0, f1, t0, t1 = event
                if f0 < 0 or f1 > 100 or t0 < 0 or t1 > args.horizon:
                    new_boundary_errors.append({"id": identifier, "event": event})
            new_named.append((identifier, plan))
            new_cells.update(cells(events))
        except (TypeError, ValueError) as exc:
            template_errors.append({"sequence": sequence, "error": str(exc)})

    all_named = q2_named + new_named
    all_conflicts = pair_conflicts(all_named)
    discrete_owners: defaultdict[tuple[int, int], list[str]] = defaultdict(list)
    for identifier, plan in all_named:
        for cell in cells(expand_plan(plan)):
            discrete_owners[cell].append(identifier)
    discrete_conflicts = sorted(
        {
            tuple(sorted((owners[left], owners[right])))
            for owners in discrete_owners.values()
            if len(owners) >= 2
            for left in range(len(owners))
            for right in range(left + 1, len(owners))
        }
    )
    duplicate_cells = [
        {"cell": list(cell), "owners": owners}
        for cell, owners in discrete_owners.items()
        if len(owners) > 1
    ]
    new_vs_q2_conflicts = [
        pair for pair in all_conflicts if pair[0].startswith("N") or pair[1].startswith("N")
    ]
    new_new_conflicts = [
        pair for pair in all_conflicts if pair[0].startswith("N") and pair[1].startswith("N")
    ]

    workbook_check = check_workbook(args.workbook, expected_rows)
    if not workbook_check["passed"]:
        errors.extend(workbook_check["errors"])
    if q2_boundary_errors:
        errors.append("fixed Q2 background has boundary errors")
    if q2_conflicts:
        errors.append("fixed Q2 background has continuous conflicts")
    errors.extend(sequence_errors)
    if template_errors:
        errors.append("new Q3 schedule has template errors")
    errors.extend(
        ["new Q3 schedule has boundary errors"] if new_boundary_errors else []
    )
    errors.extend(
        ["new Q3 schedule has continuous conflicts"] if new_vs_q2_conflicts else []
    )
    errors.extend(
        ["new Q3 schedule has duplicate discrete cells"] if duplicate_cells else []
    )

    solver = q3_result.get("solver", {})
    selected_count = len(selected)
    solver_optimality_reported = (
        q3_result.get("status") == "OPTIMAL"
        and q3_result.get("optimality") == "PROVED"
        and solver.get("status_code") == 0
        and solver.get("selected_count") == selected_count
        and solver.get("upper_bound_from_solver") == selected_count
        and float(solver.get("mip_gap", 1.0)) == 0.0
    )
    if not solver_optimality_reported:
        errors.append("Q3 solver optimality evidence is incomplete")

    payload = {
        "question": "D-Q3",
        "status": "PASS" if not errors else "FAIL",
        "result_status": q3_result.get("status"),
        "result_optimality": q3_result.get("optimality"),
        "horizon": args.horizon,
        "template": {
            "frequency_width": C_WIDTH,
            "duration": C_DURATION,
            "gap": C_GAP,
            "uses": C_USES,
            "step": C_STEP,
            "resource_cells_per_plan": C_WIDTH * C_DURATION * C_USES,
        },
        "counts": {
            "raw_plan_count": len(raw_plans),
            "fixed_q2_active_plan_count": len(q2_active),
            "fixed_q2_occupied_cell_count": len(q2_cells),
            "new_plan_count": selected_count,
            "new_occupied_cell_count": len(new_cells),
            "combined_occupied_cell_count": len(discrete_owners),
            "combined_active_plan_count": len(all_named),
            "free_cell_area_upper_bound": (args.horizon * 100 - len(q2_cells)) // (C_WIDTH * C_DURATION * C_USES),
        },
        "checks": {
            "fixed_q2_action_reconstruction": not q2_action_errors,
            "fixed_q2_boundary_errors": q2_boundary_errors,
            "fixed_q2_continuous_conflict_pairs": q2_conflicts,
            "new_plan_sequence": not sequence_errors,
            "new_template_errors": template_errors,
            "new_boundary_errors": new_boundary_errors,
            "new_vs_q2_continuous_conflict_pairs": new_vs_q2_conflicts,
            "new_new_continuous_conflict_pairs": new_new_conflicts,
            "all_continuous_conflict_pairs": all_conflicts,
            "all_discrete_conflict_pairs": discrete_conflicts,
            "continuous_equals_discrete": all_conflicts == discrete_conflicts,
            "duplicate_occupied_cells": duplicate_cells,
            "workbook": workbook_check,
            "solver_optimality_reported": solver_optimality_reported,
        },
        "errors": errors,
        "claim_boundary": (
            f"固定当前 Q2 主方案和 [0,{args.horizon})×[0,100) 资源域时，求解器状态与上下界均支持最多新增 {selected_count} 台。"
            if not errors
            else "验证未通过，不能将结果写成最多新增数量。"
        ),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "errors": len(errors), "new_plan_count": selected_count}, ensure_ascii=False))


if __name__ == "__main__":
    main()
