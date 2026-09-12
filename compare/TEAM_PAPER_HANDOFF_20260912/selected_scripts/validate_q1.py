#!/usr/bin/env python3
"""Independent audit of the formal D-Q1 conflict-pair output."""
from __future__ import annotations

from collections import Counter, defaultdict
from itertools import combinations
import csv
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook


INTERVAL_RE = re.compile(r"\s*\[\s*(-?\d+)\s*,\s*(-?\d+)\s*\)\s*")


def parse_interval(value: str) -> tuple[int, int]:
    match = INTERVAL_RE.fullmatch(value)
    if not match:
        raise ValueError(f"invalid interval: {value!r}")
    return tuple(map(int, match.groups()))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def main() -> None:
    input_path = Path("data/raw/D题/附件/附件1.xlsx")
    result_path = Path("outputs/q1/results.json")
    csv_path = Path("outputs/q1/conflict_pairs.csv")
    output_path = Path("outputs/q1/validation.json")

    workbook = load_workbook(input_path, data_only=False, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    plans: dict[str, dict] = {}
    cell_owners: defaultdict[tuple[int, int], set[str]] = defaultdict(set)
    event_rectangles: dict[str, list[tuple[int, int, int, int]]] = {}
    weighted_cell_count = 0
    event_count = 0
    for identifier, f_text, t_text, gap, uses in sheet.iter_rows(min_row=2, values_only=True):
        f0, f1 = parse_interval(f_text)
        t0, t1 = parse_interval(t_text)
        duration = t1 - t0
        rectangles = []
        for repeat in range(uses):
            start = t0 + repeat * (duration + gap)
            end = t1 + repeat * (duration + gap)
            rectangles.append((f0, f1, start, end))
            event_count += 1
            for time in range(start, end):
                for frequency in range(f0, f1):
                    cell_owners[(time, frequency)].add(identifier)
                    weighted_cell_count += 1
        plans[identifier] = {"frequency": (f0, f1), "events": rectangles}
        event_rectangles[identifier] = rectangles

    discovered_pairs: set[tuple[str, str]] = set()
    for owners in cell_owners.values():
        discovered_pairs.update(combinations(sorted(owners), 2))

    event_overlap_count = 0
    for left, right in combinations(sorted(plans), 2):
        lf0, lf1 = plans[left]["frequency"]
        rf0, rf1 = plans[right]["frequency"]
        if max(lf0, rf0) >= min(lf1, rf1):
            continue
        for _, _, lt0, lt1 in event_rectangles[left]:
            for _, _, rt0, rt1 in event_rectangles[right]:
                if max(lt0, rt0) < min(lt1, rt1):
                    event_overlap_count += 1

    result = json.loads(result_path.read_text(encoding="utf-8"))
    reported_pairs = {
        tuple(sorted((row["冲突装备1"], row["冲突装备2"])))
        for row in result["conflict_pairs"]
    }
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        csv_rows = list(csv.DictReader(handle))
    csv_pairs = {(row["冲突装备1"], row["冲突装备2"]) for row in csv_rows}

    category_counts = Counter("-".join(sorted((left[0], right[0]))) for left, right in discovered_pairs)
    all_events = [event for events in event_rectangles.values() for event in events]
    checks = {
        "plan_count_150": len(plans) == 150,
        "expanded_event_count_1300": event_count == 1300,
        "A001_second_event_pdf_example": event_rectangles["A001"][1] == (80, 90, 100, 105),
        "C083_last_event_boundary": event_rectangles["C083"][-1] == (71, 74, 641, 643),
        "time_envelope_0_643": [min(e[2] for e in all_events), max(e[3] for e in all_events)] == [0, 643],
        "frequency_envelope_0_100": [min(e[0] for e in all_events), max(e[1] for e in all_events)] == [0, 100],
        "cell_pairs_match_results_json": discovered_pairs == reported_pairs,
        "csv_pairs_match_results_json": csv_pairs == reported_pairs,
        "reported_rows_unique": len(result["conflict_pairs"]) == len(reported_pairs),
        "reported_sequence_contiguous": [int(row["序号"]) for row in csv_rows] == list(range(1, len(csv_rows) + 1)),
        "event_overlap_count_matches": event_overlap_count == result["event_overlap_count"],
        "category_counts_match": dict(sorted(category_counts.items())) == result["conflict_pairs_by_category"],
    }
    audit = {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "method": "independent occupied-cell inverted index plus event-pair recount",
        "input_sha256": sha256(input_path),
        "results_sha256": sha256(result_path),
        "plan_count": len(plans),
        "candidate_plan_pair_count": len(plans) * (len(plans) - 1) // 2,
        "expanded_event_count": event_count,
        "conflict_pair_count": len(discovered_pairs),
        "event_overlap_count": event_overlap_count,
        "weighted_occupied_cell_count": weighted_cell_count,
        "unique_occupied_cell_count": len(cell_owners),
        "max_cell_multiplicity": max(map(len, cell_owners.values())),
        "checks": checks,
    }
    output_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    if audit["status"] != "PASS":
        failed = [name for name, passed in checks.items() if not passed]
        raise AssertionError(f"Q1 independent validation failed: {failed}")
    print(json.dumps(audit, ensure_ascii=False))


if __name__ == "__main__":
    main()
