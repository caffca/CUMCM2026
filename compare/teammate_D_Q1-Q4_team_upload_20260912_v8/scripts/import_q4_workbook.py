#!/usr/bin/env python3
"""Normalize the sparse official result4.xlsx template into Q4 action JSON."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from scripts.solve_q1 import read_plans
from scripts.validate_q4 import apply_actions, category_counts, recompute_objectives


INTERVAL_RE = re.compile(r"\s*\[\s*(-?\d+)\s*,\s*(-?\d+)\s*\)\s*")


def parse_interval(value: object, label: str, identifier: str) -> tuple[int, int]:
    match = INTERVAL_RE.fullmatch(str(value or ""))
    if not match:
        raise ValueError(f"{identifier}: invalid {label} interval {value!r}")
    left, right = map(int, match.groups())
    if right <= left:
        raise ValueError(f"{identifier}: non-positive {label} interval {value!r}")
    return left, right


def interval_text(interval: tuple[int, int] | None) -> str | None:
    return None if interval is None else f"[{interval[0]},{interval[1]})"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data/raw/D题/附件/附件1.xlsx"))
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    plans = read_plans(args.input)
    by_id = {plan["id"]: plan for plan in plans}
    sheet = load_workbook(args.workbook, data_only=True, read_only=True).active
    actions = []
    seen = set()
    errors = []
    for row in sheet.iter_rows(min_row=2, max_col=5, values_only=True):
        identifier, frequency_value, time_value, gap_value, revoke_value = row
        if identifier in (None, ""):
            continue
        identifier = str(identifier).strip()
        if identifier in seen:
            errors.append(f"duplicate row for {identifier}")
            continue
        seen.add(identifier)
        plan = by_id.get(identifier)
        if plan is None:
            errors.append(f"unknown equipment id {identifier}")
            continue
        has_frequency = frequency_value not in (None, "")
        has_time = time_value not in (None, "")
        has_gap = gap_value not in (None, "")
        has_revoke = revoke_value not in (None, "")
        if has_revoke:
            if has_frequency or has_time or has_gap or str(revoke_value).strip() != "是":
                errors.append(f"invalid revoke row for {identifier}")
                continue
            actions.append({
                "装备编号": identifier,
                "类别": plan["category"],
                "动作": "revoke",
                "频段区间": None,
                "时间区间": None,
                "调整后间隔": None,
                "频移": 0,
                "时移": 0,
                "间隔变化": 0,
                "撤销": True,
            })
            continue
        if sum([has_frequency, has_time, has_gap]) != 1:
            errors.append(f"expected exactly one populated action field for {identifier}")
            continue
        if has_frequency:
            frequency = parse_interval(frequency_value, "frequency", identifier)
            df = frequency[0] - plan["frequency"][0]
            if df == 0 or frequency[1] - frequency[0] != plan["frequency"][1] - plan["frequency"][0] or abs(df) > 10:
                errors.append(f"invalid frequency adjustment for {identifier}")
                continue
            actions.append({
                "装备编号": identifier,
                "类别": plan["category"],
                "动作": "freq",
                "频段区间": interval_text(frequency),
                "时间区间": None,
                "调整后间隔": plan["gap"],
                "频移": df,
                "时移": 0,
                "间隔变化": 0,
                "撤销": False,
            })
        elif has_time:
            time = parse_interval(time_value, "time", identifier)
            dt = time[0] - plan["time"][0]
            if dt == 0 or time[1] - time[0] != plan["time"][1] - plan["time"][0] or abs(dt) > 5:
                errors.append(f"invalid time adjustment for {identifier}")
                continue
            actions.append({
                "装备编号": identifier,
                "类别": plan["category"],
                "动作": "time",
                "频段区间": None,
                "时间区间": interval_text(time),
                "调整后间隔": plan["gap"],
                "频移": 0,
                "时移": dt,
                "间隔变化": 0,
                "撤销": False,
            })
        else:
            if plan["category"] != "C" or not isinstance(gap_value, (int, float)) or int(gap_value) != gap_value:
                errors.append(f"invalid gap adjustment for {identifier}")
                continue
            new_gap = int(gap_value)
            dg = new_gap - plan["gap"]
            if dg == 0 or abs(dg) > 10 or new_gap < 0:
                errors.append(f"gap limit violated for {identifier}")
                continue
            actions.append({
                "装备编号": identifier,
                "类别": plan["category"],
                "动作": "gap",
                "频段区间": None,
                "时间区间": None,
                "调整后间隔": new_gap,
                "频移": 0,
                "时移": 0,
                "间隔变化": dg,
                "撤销": False,
            })

    if errors:
        raise ValueError("; ".join(errors))
    adjusted, action_errors = apply_actions(plans, actions)
    if action_errors:
        raise ValueError("; ".join(action_errors))
    objectives = recompute_objectives(adjusted, plans)
    counts = category_counts(adjusted)
    payload = {
        "question": "D-Q4",
        "status": "IMPORTED_FEASIBLE_CANDIDATE",
        "source_workbook": str(args.workbook),
        "plan_count": len(plans),
        "summary": {
            "plans": len(plans),
            "selected_candidates": len(plans),
            "actions": actions,
            "counts_by_category": counts,
        },
        "optimization": {"objective_values": objectives},
        "import_checks": {
            "workbook_action_rows": len(seen),
            "canonical_action_rows": len(actions),
            "implicit_keep_rows": len(plans) - len(actions),
            "exactly_one_action_field": True,
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "action_rows": len(actions), "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
