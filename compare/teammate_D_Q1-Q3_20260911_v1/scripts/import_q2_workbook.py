#!/usr/bin/env python3
"""Normalize a result2 workbook into Q2 action JSON.

Both the dense template (final frequency and time on every non-revoked action
row) and the sparse template (only the changed parameter is filled) are
accepted. The importer compares the supplied intervals with the source plan,
infers the single changed parameter, and rejects rows that change both
parameters or violate move limits. Omitted rows and rows with no explicit
action are treated as implicit keep decisions. The result is still an
incumbent until validate_q2.py independently checks all expanded events and
resource cells.
"""
from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path
import re

from openpyxl import load_workbook

from scripts.solve_q1 import read_plans


INTERVAL_RE = re.compile(r"\s*\[\s*(-?\d+)\s*,\s*(-?\d+)\s*\)\s*")


def parse_interval(value: object, label: str, identifier: str) -> tuple[int, int]:
    match = INTERVAL_RE.fullmatch(str(value or ""))
    if not match:
        raise ValueError(f"{identifier}: invalid {label} interval {value!r}")
    left, right = map(int, match.groups())
    if right <= left:
        raise ValueError(f"{identifier}: non-positive {label} interval {value!r}")
    return left, right


def interval_text(interval: tuple[int, int]) -> str:
    return f"[{interval[0]},{interval[1]})"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data/raw/D题/附件/附件1.xlsx"))
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--source-label", default="externally supplied result2.xlsx")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    plans = read_plans(args.input)
    by_id = {plan["id"]: plan for plan in plans}
    sheet = load_workbook(args.workbook, data_only=True, read_only=True).active
    actions: list[dict] = []
    seen: set[str] = set()
    errors: list[str] = []

    for values in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
        identifier, frequency_value, time_value, revoke_value = values
        if identifier is None:
            continue
        identifier = str(identifier).strip()
        if identifier in seen:
            errors.append(f"duplicate row for {identifier}")
            continue
        seen.add(identifier)
        plan = by_id.get(identifier)
        if plan is None:
            errors.append(f"unknown equipment id {identifier}")
            continue

        has_frequency = frequency_value not in (None, "")
        has_time = time_value not in (None, "")
        has_revoke = revoke_value not in (None, "")

        if has_revoke:
            if has_frequency or has_time:
                errors.append(f"{identifier}: revoke row cannot also contain frequency/time")
                continue
            actions.append({
                "装备编号": identifier,
                "类别": plan["category"],
                "动作": "revoke",
                "频段区间": None,
                "时间区间": None,
                "频移": 0,
                "时移": 0,
                "撤销": True,
            })
            continue

        # A row with no explicit action is an implicit keep row. This is
        # useful when a workbook lists all 150 plans instead of only the
        # changed/revoked rows.
        if not has_frequency and not has_time:
            continue

        try:
            frequency = (
                parse_interval(frequency_value, "frequency", identifier)
                if has_frequency
                else plan["frequency"]
            )
            time = (
                parse_interval(time_value, "time", identifier)
                if has_time
                else plan["time"]
            )
        except ValueError as exc:
            errors.append(str(exc))
            continue
        df = frequency[0] - plan["frequency"][0]
        dt = time[0] - plan["time"][0]
        same_frequency_width = frequency[1] - frequency[0] == plan["frequency"][1] - plan["frequency"][0]
        same_time_width = time[1] - time[0] == plan["time"][1] - plan["time"][0]
        if not same_frequency_width or not same_time_width:
            errors.append(f"{identifier}: interval width/duration changed")
            continue
        if df == 0 and dt == 0:
            errors.append(f"{identifier}: explicit row does not change either parameter")
            continue
        if df != 0 and dt != 0:
            errors.append(f"{identifier}: expected exactly one non-zero shift, got df={df}, dt={dt}")
            continue
        action = "freq" if df else "time"
        if abs(df) > 10 or abs(dt) > 5:
            errors.append(f"{identifier}: move limit exceeded, df={df}, dt={dt}")
            continue
        actions.append({
            "装备编号": identifier,
            "类别": plan["category"],
            "动作": action,
            "频段区间": interval_text(frequency),
            "时间区间": interval_text(time),
            "频移": df,
            "时移": dt,
            "撤销": False,
        })

    if errors:
        raise ValueError("; ".join(errors))

    action_counts = Counter(row["动作"] for row in actions)
    category_counts = {
        category: {
            "保留数量": sum(plan["category"] == category for plan in plans) - sum(row["类别"] == category for row in actions),
            "调整数量": sum(row["类别"] == category and row["动作"] != "revoke" for row in actions),
            "撤销数量": sum(row["类别"] == category and row["动作"] == "revoke" for row in actions),
        }
        for category in "ABC"
    }
    payload = {
        "question": "D-Q2",
        "status": "IMPORTED_FEASIBLE_CANDIDATE",
        "optimality": "UNVERIFIED_BY_IMPORTER",
        "source_workbook": args.source_label,
        "plan_count": len(plans),
        "objective_values": {
            "revocations": action_counts["revoke"],
            "revoke_A": category_counts["A"]["撤销数量"],
            "revoke_B": category_counts["B"]["撤销数量"],
            "adjusted_plans": action_counts["freq"] + action_counts["time"],
            "adjust_A": category_counts["A"]["调整数量"],
            "adjust_B": category_counts["B"]["调整数量"],
            "frequency_shift_sum": sum(abs(row["频移"]) for row in actions),
            "time_shift_sum": sum(abs(row["时移"]) for row in actions),
            "normalized_shift_score_x10": sum(abs(row["频移"]) + 2 * abs(row["时移"]) for row in actions),
        },
        "summary": {
            "plans": len(plans),
            "selected_candidates": len(plans),
            "actions": actions,
            "counts_by_category": category_counts,
        },
        "import_checks": {
            "workbook_action_rows": len(seen),
            "canonical_action_rows": len(actions),
            "implicit_keep_rows": len(plans) - len(actions),
            "exactly_one_parameter_changed_on_adjustment": True,
            "move_limits_checked": True,
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), **payload["objective_values"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
