#!/usr/bin/env python3
"""Check that result2.xlsx exactly represents the selected Q2 JSON actions."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


HEADERS = ["用频装备编号", "调整后频段区间", "调整后时间区间", "是否撤销用频计划"]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--result", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    payload = json.loads(args.result.read_text(encoding="utf-8"))
    expected = {}
    for action in payload["summary"]["actions"]:
        identifier = action["装备编号"]
        if identifier in expected:
            raise ValueError(f"duplicate JSON action {identifier}")
        expected[identifier] = action

    workbook = load_workbook(args.workbook, read_only=True, data_only=False)
    sheet = workbook.active
    actual_headers = [cell.value for cell in sheet[1]]
    errors = []
    if actual_headers != HEADERS:
        errors.append(f"headers mismatch: {actual_headers!r}")
    actual = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None or value == "" for value in row):
            continue
        identifier, frequency, time, revoke = row
        if identifier in actual:
            errors.append(f"duplicate workbook action {identifier}")
            continue
        nonempty = sum(value not in (None, "") for value in (frequency, time, revoke))
        if nonempty != 1:
            errors.append(f"row for {identifier} does not contain exactly one action field")
            continue
        if revoke not in (None, ""):
            action = "revoke"
        elif frequency not in (None, ""):
            action = "freq"
        else:
            action = "time"
        actual[identifier] = {
            "动作": action,
            "频段区间": frequency if action == "freq" else None,
            "时间区间": time if action == "time" else None,
        }

    for identifier in sorted(set(expected) | set(actual)):
        if identifier not in expected:
            errors.append(f"unexpected workbook action {identifier}")
            continue
        if identifier not in actual:
            errors.append(f"missing workbook action {identifier}")
            continue
        reference = expected[identifier]
        observed = actual[identifier]
        fields = ["动作"]
        if observed["动作"] == "freq":
            fields.append("频段区间")
        elif observed["动作"] == "time":
            fields.append("时间区间")
        for field in fields:
            if observed[field] != reference[field]:
                errors.append(f"{identifier} {field}: expected {reference[field]!r}, got {observed[field]!r}")

    output = {
        "status": "PASS" if not errors else "FAIL",
        "expected_action_count": len(expected),
        "actual_action_count": len(actual),
        "errors": errors,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": output["status"], "errors": len(errors), "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
