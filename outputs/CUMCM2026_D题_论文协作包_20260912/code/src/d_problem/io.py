"""Input and Q1 result-template adapters."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd

from .domain import ConflictPair, Plan
from .candidates import CandidateState


_INTERVAL_RE = re.compile(r"^\s*\[\s*(-?\d+)\s*,\s*(-?\d+)\s*\)\s*$")
REQUIRED_COLUMNS = ("用频装备编号", "频段区间", "时间区间", "间隔时长", "使用次数")


def parse_half_open_interval(value: object) -> tuple[int, int]:
    match = _INTERVAL_RE.match(str(value))
    if match is None:
        raise ValueError(f"invalid half-open interval: {value!r}")
    left, right = (int(x) for x in match.groups())
    if right <= left:
        raise ValueError(f"interval must have positive length: {value!r}")
    return left, right


def load_plans(path: str | Path) -> list[Plan]:
    frame = pd.read_excel(path, dtype={"用频装备编号": str})
    missing = [column for column in REQUIRED_COLUMNS if column not in frame.columns]
    if missing:
        raise ValueError(f"missing required columns: {missing}")

    plans: list[Plan] = []
    seen: set[str] = set()
    for row_number, row in frame.iterrows():
        plan_id = str(row["用频装备编号"]).strip()
        if not plan_id or plan_id == "nan":
            raise ValueError(f"empty plan id at Excel data row {row_number + 2}")
        if plan_id in seen:
            raise ValueError(f"duplicate plan id: {plan_id}")
        seen.add(plan_id)
        f_start, f_end = parse_half_open_interval(row["频段区间"])
        t_start, t_end = parse_half_open_interval(row["时间区间"])
        gap = int(row["间隔时长"])
        count = int(row["使用次数"])
        if gap < 0 or count < 1:
            raise ValueError(f"invalid gap/count for {plan_id}: {gap}, {count}")
        if not plan_id[0] in {"A", "B", "C"}:
            raise ValueError(f"unknown category in plan id: {plan_id}")
        plans.append(
            Plan(
                plan_id=plan_id,
                category=plan_id[0],
                f_start=f_start,
                f_end=f_end,
                t_start=t_start,
                t_end=t_end,
                gap=gap,
                count=count,
            )
        )
    return plans


def write_q1_result(
    output_path: str | Path,
    conflicts: Iterable[ConflictPair],
    template_path: str | Path | None = None,
) -> None:
    """Write Q1 rows while leaving the supplied blank template untouched."""

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if template_path is not None:
        workbook = openpyxl.load_workbook(template_path)
    else:
        workbook = openpyxl.Workbook()
        workbook.active.append(["序号", "冲突装备1", "冲突设备2"])
    sheet = workbook[workbook.sheetnames[0]]

    # Preserve the header and discard only rows in the newly-created output.
    # The caller never passes the raw template as output_path.
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for index, pair in enumerate(conflicts, start=1):
        sheet.append([index, pair.plan_i, pair.plan_j])
    workbook.save(output)


def _format_interval(left: int, right: int) -> str:
    return f"[{left},{right})"


def write_q2_result(
    output_path: str | Path,
    selected: Iterable[CandidateState],
    template_path: str | Path | None = None,
) -> None:
    """Write only adjusted/cancelled plans to the result2 template."""

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    if template_path is not None:
        workbook = openpyxl.load_workbook(template_path)
    else:
        workbook = openpyxl.Workbook()
        workbook.active.append(
            ["用频装备编号", "调整后频段区间", "调整后时间区间", "是否撤销用频计划"]
        )
    sheet = workbook[workbook.sheetnames[0]]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    for candidate in selected:
        if candidate.action == "keep":
            continue
        plan = candidate.plan
        if candidate.action == "cancel":
            sheet.append([plan.plan_id, None, None, "是"])
            continue
        f_start = plan.f_start + candidate.state.f_shift
        f_end = plan.f_end + candidate.state.f_shift
        t_start = plan.t_start + candidate.state.t_shift
        t_end = plan.t_end + candidate.state.t_shift
        sheet.append(
            [
                plan.plan_id,
                _format_interval(f_start, f_end),
                _format_interval(t_start, t_end),
                "否",
            ]
        )
    workbook.save(output)
