"""Validate that outputs/q4/result4.xlsx matches a selected Q4 snapshot.

This checks the template adapter, not the optimization itself.  The selected
state JSON is still validated by the canonical schedule validator in the solve
report; this script makes sure the user-facing workbook did not lose a shift,
gap, or cancellation while being written.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from d_problem.io import load_plans  # noqa: E402


INTERVAL_RE = re.compile(r"^\[(-?\d+),(-?\d+)\)$")


def parse_interval(value: object) -> tuple[int, int] | None:
    if value is None:
        return None
    match = INTERVAL_RE.fullmatch(str(value))
    if match is None:
        raise ValueError(f"invalid workbook interval: {value!r}")
    return int(match.group(1)), int(match.group(2))


def validate(workbook_path: Path, selected_path: Path, input_path: Path) -> dict[str, object]:
    selected_payload = json.loads(selected_path.read_text(encoding="utf-8"))
    selected = selected_payload.get("selected", [])
    plans = {plan.plan_id: plan for plan in load_plans(input_path)}
    expected = {
        record["plan_id"]: record
        for record in selected
        if record.get("action") != "keep"
    }
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)
    sheet = workbook.active
    headers = [sheet.cell(1, column).value for column in range(1, 6)]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    rows = [row for row in rows if any(value is not None for value in row)]
    mismatches: list[str] = []
    seen: set[str] = set()
    for row in rows:
        plan_id, frequency, time_interval, gap, cancelled = row[:5]
        plan_id = str(plan_id).strip()
        seen.add(plan_id)
        if plan_id not in expected:
            mismatches.append(f"unexpected plan {plan_id}")
            continue
        record = expected[plan_id]
        action = record["action"]
        if action == "cancel":
            if not (frequency is None and time_interval is None and gap is None and cancelled == "是"):
                mismatches.append(f"cancel row mismatch: {plan_id}")
            continue
        if plan_id not in plans:
            mismatches.append(f"unknown plan {plan_id}")
            continue
        plan = plans[plan_id]
        expected_frequency = (
            plan.f_start + int(record["f_shift"]),
            plan.f_end + int(record["f_shift"]),
        )
        expected_time = (
            plan.t_start + int(record["t_shift"]),
            plan.t_end + int(record["t_shift"]),
        )
        if parse_interval(frequency) != expected_frequency:
            mismatches.append(f"frequency mismatch: {plan_id}")
        if parse_interval(time_interval) != expected_time:
            mismatches.append(f"time mismatch: {plan_id}")
        if int(gap) != int(record["adjusted_gap"]) or cancelled != "否":
            mismatches.append(f"gap/cancel mismatch: {plan_id}")
    missing = sorted(set(expected) - seen)
    if missing:
        mismatches.append(f"missing plans: {','.join(missing)}")
    return {
        "workbook": str(workbook_path),
        "selected": str(selected_path),
        "headers": headers,
        "nonempty_rows": len(rows),
        "expected_changed_rows": len(expected),
        "seen_plan_count": len(seen),
        "cancelled_rows": sum(row[4] == "是" for row in rows),
        "mismatch_count": len(mismatches),
        "mismatches": mismatches,
        "ok": not mismatches and len(rows) == len(expected),
    }


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=root / "outputs/q4/result4.xlsx")
    parser.add_argument("--selected", type=Path, default=root / "outputs/q4/selected_T.json")
    parser.add_argument("--input", type=Path, default=root / "D题/附件/附件1.xlsx")
    parser.add_argument("--output", type=Path, default=root / "outputs/q4/result4_readback_check.json")
    args = parser.parse_args()
    report = validate(args.workbook, args.selected, args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    if not report["ok"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
