"""Solve CUMCM2026 D, Q3 for a frozen Q2 schedule.

The Q3 decision is a maximum set-packing problem.  A new C-class plan is
identified by its initial frequency start and initial time start; its 12
periodic occurrences are then fixed by the C template.  Candidates that
intersect the frozen Q2 occupancy are removed first.  For the remaining
candidates, every integer time-frequency cell is an at-most-one resource
clique.  This is equivalent to the half-open rectangle conflict definition
because all endpoints are integer grid points.
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from ortools.sat.python import cp_model

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from d_problem.conflicts import expand_occurrences  # noqa: E402
from d_problem.domain import Plan, PlanState  # noqa: E402
from d_problem.io import load_plans  # noqa: E402
from d_problem.validation import validate_schedule  # noqa: E402


@dataclass(frozen=True, slots=True)
class NewCandidate:
    index: int
    f_start: int
    t_start: int
    f_end: int
    t_end: int
    cells: tuple[tuple[int, int], ...]


def _interval(left: int, right: int) -> str:
    return f"[{left},{right})"


def _load_q2_states(path: Path, plans: list[Plan]) -> list[PlanState]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    records = payload.get("selected")
    if not isinstance(records, list):
        raise ValueError(f"Q2 selection file has no selected list: {path}")
    by_id = {plan.plan_id: plan for plan in plans}
    states: list[PlanState] = []
    seen: set[str] = set()
    for record in records:
        plan_id = str(record["plan_id"])
        if plan_id not in by_id:
            raise ValueError(f"unknown Q2 plan id {plan_id}")
        if plan_id in seen:
            raise ValueError(f"duplicate Q2 plan id {plan_id}")
        seen.add(plan_id)
        action = str(record.get("action", "keep"))
        if action not in {"keep", "frequency", "time", "cancel"}:
            raise ValueError(f"unsupported Q2 action {action!r} for {plan_id}")
        states.append(
            PlanState(
                by_id[plan_id],
                action=action,
                f_shift=int(record.get("f_shift", 0)),
                t_shift=int(record.get("t_shift", 0)),
            )
        )
    missing = sorted(set(by_id) - seen)
    if missing:
        raise ValueError(f"Q2 selection is not complete; missing {missing[:5]}")
    return states


def _c_template(plans: list[Plan]) -> dict[str, int]:
    c_plans = [plan for plan in plans if plan.category == "C"]
    if not c_plans:
        raise ValueError("附件1 contains no C-class plan")
    signatures = Counter(
        (plan.f_end - plan.f_start, plan.t_end - plan.t_start, plan.gap, plan.count)
        for plan in c_plans
    )
    signature, count = signatures.most_common(1)[0]
    if count != len(c_plans):
        raise ValueError(
            "C-class template is not homogeneous; pass an explicit template "
            "or resolve the data semantics before Q3"
        )
    width, duration, gap, uses = signature
    return {"width": width, "duration": duration, "gap": gap, "count": uses}


def _cells(
    f_start: int,
    t_start: int,
    width: int,
    duration: int,
    gap: int,
    count: int,
) -> tuple[tuple[int, int], ...]:
    period = duration + gap
    return tuple(
        (frequency, t_start + k * period + offset)
        for k in range(count)
        for frequency in range(f_start, f_start + width)
        for offset in range(duration)
    )


def _generate_candidates(
    occupied: set[tuple[int, int]],
    horizon: int,
    template: dict[str, int],
) -> tuple[list[NewCandidate], dict[str, int]]:
    width = template["width"]
    duration = template["duration"]
    gap = template["gap"]
    count = template["count"]
    period = duration + gap
    span = duration + (count - 1) * period
    max_f = 100 - width
    max_t = horizon - span
    if max_f < 0 or max_t < 0:
        raise ValueError("C template cannot fit in the declared resource domain")

    candidates: list[NewCandidate] = []
    raw = 0
    rejected = 0
    for f_start in range(max_f + 1):
        for t_start in range(max_t + 1):
            raw += 1
            cells = _cells(f_start, t_start, width, duration, gap, count)
            if any(cell in occupied for cell in cells):
                rejected += 1
                continue
            index = len(candidates)
            candidates.append(
                NewCandidate(
                    index=index,
                    f_start=f_start,
                    t_start=t_start,
                    f_end=f_start + width,
                    t_end=t_start + duration,
                    cells=cells,
                )
            )
    return candidates, {
        "raw_candidate_count": raw,
        "q2_conflict_rejected": rejected,
        "compatible_candidate_count": len(candidates),
        "frequency_start_min": 0,
        "frequency_start_max": max_f,
        "time_start_min": 0,
        "time_start_max": max_t,
        "period": period,
        "span": span,
    }


def _greedy_hint(candidates: list[NewCandidate], cell_map: dict[tuple[int, int], list[int]]) -> set[int]:
    """Deterministic maximal packing used only as a CP-SAT warm start."""

    degrees = {
        candidate.index: sum(len(cell_map[cell]) - 1 for cell in candidate.cells)
        for candidate in candidates
    }
    used: set[tuple[int, int]] = set()
    chosen: set[int] = set()
    order = sorted(candidates, key=lambda item: (degrees[item.index], item.t_start, item.f_start))
    for candidate in order:
        if any(cell in used for cell in candidate.cells):
            continue
        chosen.add(candidate.index)
        used.update(candidate.cells)
    return chosen


def _solve(candidates: list[NewCandidate], time_limit: float, workers: int, seed: int):
    cell_map: dict[tuple[int, int], list[int]] = defaultdict(list)
    for candidate in candidates:
        for cell in candidate.cells:
            cell_map[cell].append(candidate.index)
    cliques = [tuple(indices) for indices in cell_map.values() if len(indices) > 1]

    model = cp_model.CpModel()
    variables = [model.NewBoolVar(f"z_{candidate.index}") for candidate in candidates]
    for clique in cliques:
        model.AddAtMostOne(variables[index] for index in clique)
    objective = sum(variables)
    model.Maximize(objective)

    hint = _greedy_hint(candidates, cell_map)
    for candidate in candidates:
        model.add_hint(variables[candidate.index], int(candidate.index in hint))

    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = time_limit
    solver.parameters.num_search_workers = workers
    solver.parameters.random_seed = seed
    solver.parameters.log_search_progress = False
    started = time.perf_counter()
    status = solver.Solve(model)
    elapsed = time.perf_counter() - started
    status_name = solver.StatusName(status)
    selected: list[NewCandidate] = []
    if status in (cp_model.OPTIMAL, cp_model.FEASIBLE):
        selected = [candidate for candidate in candidates if solver.Value(variables[candidate.index])]
    value = len(selected) if selected else None
    bound = float(solver.BestObjectiveBound()) if status in (cp_model.OPTIMAL, cp_model.FEASIBLE) else None
    gap = None if value is None or bound is None else (bound - value) / max(1.0, abs(value))
    return selected, {
        "status": status_name,
        "objective_value": value,
        "best_bound": bound,
        "gap": gap,
        "seconds": elapsed,
        "time_limit_seconds": time_limit,
        "workers": workers,
        "seed": seed,
        "clique_count": len(cliques),
        "occupied_cell_count": len(cell_map),
        "greedy_hint_value": len(hint),
    }


def _write_result(path: Path, template_path: Path, selected: list[NewCandidate], first_id: int) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(template_path)
    sheet = workbook[workbook.sheetnames[0]]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for offset, candidate in enumerate(sorted(selected, key=lambda item: (item.t_start, item.f_start)), start=0):
        sheet.append(
            [
                f"C{first_id + offset:03d}",
                _interval(candidate.f_start, candidate.f_end),
                _interval(candidate.t_start, candidate.t_end),
            ]
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _write_summary(path: Path, report: dict[str, object]) -> None:
    solve = report["solver"]
    validation = report["validation"]
    counts = report["candidate_counts"]
    lines = [
        "# Q3 求解摘要",
        "",
        f"- Q2 input: `{report['q2_input']}`",
        f"- horizon: `[0,{report['horizon']})`",
        f"- C template: `{report['template']}`",
        "",
        "## 候选与求解",
        "",
        f"- raw candidates: `{counts['raw_candidate_count']}`",
        f"- rejected by frozen Q2: `{counts['q2_conflict_rejected']}`",
        f"- compatible candidates: `{counts['compatible_candidate_count']}`",
        f"- resource-cell cliques: `{solve['clique_count']}`",
        f"- solver status: `{solve['status']}`",
        f"- selected new C plans: `{solve['objective_value']}`",
        f"- best bound: `{solve['best_bound']}`",
        f"- gap: `{solve['gap']}`",
        f"- seconds: `{solve['seconds']:.3f}`",
        "",
        "## 统一复验",
        "",
        f"- plans after merge: `{validation['plan_count']}`",
        f"- occurrence count: `{validation['occurrence_count']}`",
        f"- conflict count: `{validation['conflict_count']}`",
        f"- boundary violations: `{validation['boundary_violations']}`",
        f"- ok: `{validation['ok']}`",
        "",
        "Q3 的最大值只针对报告所列 Q2 排程、时间域和 C 类模板成立；若 Q2 排程发生变化，应重新求解。",
        "若 solver status 不是 OPTIMAL，selected new C plans 只能称为已验证可行值，不能称为全局最大值。",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=ROOT / "D题/附件/附件1.xlsx")
    parser.add_argument("--q2-selected", type=Path, default=ROOT / "outputs/q2/Q2-T-incumbent-v1.json")
    parser.add_argument("--template", type=Path, default=ROOT / "D题/附件/附件2/result3.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/q3/result3.xlsx")
    parser.add_argument("--summary", type=Path, default=ROOT / "outputs/q3/summary.md")
    parser.add_argument("--report", type=Path, default=ROOT / "outputs/q3/solver_report.json")
    parser.add_argument("--selected", type=Path, default=ROOT / "outputs/q3/selected_Q2-T-incumbent-v1.json")
    parser.add_argument("--horizon", type=int, default=643)
    parser.add_argument("--time-limit", type=float, default=600.0)
    parser.add_argument("--workers", type=int, default=16)
    parser.add_argument("--seed", type=int, default=20260912)
    parser.add_argument("--first-id", type=int, default=91)
    args = parser.parse_args()

    plans = load_plans(args.input)
    q2_states = _load_q2_states(args.q2_selected, plans)
    q2_validation = validate_schedule(q2_states, horizon=args.horizon)
    if not q2_validation.ok:
        raise ValueError(f"frozen Q2 input is not valid: {q2_validation}")

    occupied: set[tuple[int, int]] = set()
    for state in q2_states:
        for occurrence in expand_occurrences(state):
            occupied.update(
                (frequency, time_slot)
                for frequency in range(occurrence.f_start, occurrence.f_end)
                for time_slot in range(occurrence.t_start, occurrence.t_end)
            )

    template = _c_template(plans)
    candidates, candidate_counts = _generate_candidates(occupied, args.horizon, template)
    selected, solver_report = _solve(candidates, args.time_limit, args.workers, args.seed)

    new_states: list[PlanState] = []
    selected_records: list[dict[str, object]] = []
    for offset, candidate in enumerate(sorted(selected, key=lambda item: (item.t_start, item.f_start))):
        plan = Plan(
            plan_id=f"C{args.first_id + offset:03d}",
            category="C",
            f_start=candidate.f_start,
            f_end=candidate.f_end,
            t_start=candidate.t_start,
            t_end=candidate.t_end,
            gap=template["gap"],
            count=template["count"],
        )
        new_states.append(PlanState(plan))
        selected_records.append(
            {
                "plan_id": plan.plan_id,
                "frequency_interval": _interval(plan.f_start, plan.f_end),
                "time_interval": _interval(plan.t_start, plan.t_end),
                "f_start": plan.f_start,
                "t_start": plan.t_start,
                "width": template["width"],
                "duration": template["duration"],
                "gap": template["gap"],
                "count": template["count"],
            }
        )

    merged_validation = validate_schedule(q2_states + new_states, horizon=args.horizon)
    _write_result(args.output, args.template, selected, args.first_id)
    report: dict[str, object] = {
        "question": "Q3",
        "q2_input": str(args.q2_selected),
        "input": str(args.input),
        "horizon": args.horizon,
        "template": template,
        "q2_validation": {
            "plan_count": q2_validation.plan_count,
            "occurrence_count": q2_validation.occurrence_count,
            "conflict_count": q2_validation.conflict_count,
            "boundary_violations": q2_validation.boundary_violations,
            "ok": q2_validation.ok,
        },
        "candidate_counts": candidate_counts,
        "existing_occupied_cell_count": len(occupied),
        "solver": solver_report,
        "selected_new_plans": selected_records,
        "validation": {
            "plan_count": merged_validation.plan_count,
            "occurrence_count": merged_validation.occurrence_count,
            "conflict_count": merged_validation.conflict_count,
            "boundary_violations": merged_validation.boundary_violations,
            "ok": merged_validation.ok,
        },
        "proven_optimal": solver_report["status"] == "OPTIMAL",
    }
    _write_json(args.selected, {"q2_input": str(args.q2_selected), "selected": selected_records, "validation": report["validation"]})
    _write_json(args.report, report)
    _write_summary(args.summary, report)
    print(json.dumps({"status": solver_report["status"], "new_C_count": len(selected), "conflicts": merged_validation.conflict_count, "result": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
