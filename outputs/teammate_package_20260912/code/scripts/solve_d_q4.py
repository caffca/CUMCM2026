"""Solve CUMCM2026 D, Q4 with optional C-class interval-gap states.

Q4 deliberately starts from the raw attachment-1 plans.  It is not a repair
of Q2: the original conflict instance is re-optimized after adding the legal
C-only gap action.  The exact state model and the common lexicographic
objective implementation from Q2 are reused; only candidate generation and
the result4 workbook adapter are Q4-specific.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from d_problem.io import load_plans  # noqa: E402
from d_problem.objectives import objective_names  # noqa: E402
from d_problem.q2_model import (  # noqa: E402
    SolverConfig,
    solve_lexicographic,
    validate_selected_pairwise,
)
from d_problem.q4_candidates import (  # noqa: E402
    Q4Candidate,
    generate_all_candidates,
    generate_cell_cliques,
    generate_conflict_edges,
    verify_mask_sample,
)
from d_problem.validation import validate_schedule  # noqa: E402


def _record(candidate: Q4Candidate) -> dict[str, object]:
    state = candidate.state
    return {
        "index": candidate.index,
        "state_id": candidate.state_id,
        "plan_id": candidate.plan.plan_id,
        "category": candidate.category,
        "action": candidate.action,
        "shift": candidate.shift,
        "f_shift": state.f_shift,
        "t_shift": state.t_shift,
        "gap_shift": state.gap_shift,
        "adjusted_gap": candidate.plan.gap + state.gap_shift,
        "displacement_cost": candidate.displacement_cost,
    }


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _load_hint(path: Path, flat: list[Q4Candidate]) -> dict[int, int]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    by_state = {candidate.state_id: candidate.index for candidate in flat}
    hints: dict[int, int] = {}
    for record in payload.get("selected", []):
        state_id = str(record.get("state_id", ""))
        if state_id in by_state:
            hints[by_state[state_id]] = 1
            continue
        # Older selected snapshots may not carry state_id.  Reconstruct the
        # deterministic ID from the action and the relevant shift.
        plan_id = str(record["plan_id"])
        action = str(record.get("action", "keep"))
        if action == "frequency":
            shift = int(record.get("f_shift", 0))
        elif action == "time":
            shift = int(record.get("t_shift", 0))
        elif action == "gap":
            shift = int(record.get("gap_shift", 0))
        else:
            shift = 0
        candidate_id = f"{plan_id}:{action}:{shift:+d}"
        if candidate_id in by_state:
            hints[by_state[candidate_id]] = 1
    return hints


def _interval(left: int, right: int) -> str:
    return f"[{left},{right})"


def _write_result4(
    output_path: Path,
    template_path: Path,
    selected: list[Q4Candidate],
) -> None:
    import openpyxl

    workbook = openpyxl.load_workbook(template_path)
    sheet = workbook[workbook.sheetnames[0]]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for candidate in selected:
        if candidate.action == "keep":
            continue
        plan = candidate.plan
        if candidate.action == "cancel":
            sheet.append([plan.plan_id, None, None, None, "是"])
            continue
        f_start = plan.f_start + candidate.state.f_shift
        f_end = plan.f_end + candidate.state.f_shift
        t_start = plan.t_start + candidate.state.t_shift
        t_end = plan.t_end + candidate.state.t_shift
        gap = plan.gap + candidate.state.gap_shift
        sheet.append(
            [
                plan.plan_id,
                _interval(f_start, f_end),
                _interval(t_start, t_end),
                gap,
                "否",
            ]
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _write_summary(path: Path, report: dict[str, object]) -> None:
    from collections import Counter

    metrics = report.get("metrics", {})
    lines = [
        "# Q4 求解摘要",
        "",
        "Q4 从原始 Q1 计划重新建模；C 类增加间隔调整候选，不复用某个 Q2 结果作为约束。",
        "",
        f"- horizon: `[0,{report['horizon']})`",
        f"- policy: `{report['policy']}`",
        f"- proven_optimal: `{report.get('proven_optimal', False)}`",
        f"- candidate states: `{report['candidate_summary']['candidate_state_count']}`",
        f"- state conflict edges (audit): `{report['candidate_summary']['state_conflict_edges']}`",
        f"- cell cliques (solve): `{report['candidate_summary']['cell_clique_count']}`",
        "",
        "## 目标向量",
        "",
        "| 指标 | 值 |",
        "|---|---:|",
    ]
    for name in report["objective_names"]:
        lines.append(f"| {name} | {metrics.get(name)} |")
    lines += [
        "",
        "## 求解层状态",
        "",
        "| 层 | 状态 | 值 | 下界 | gap | 秒 |",
        "|---|---|---:|---:|---:|---:|",
    ]
    for layer in report["layers"]:
        lines.append(
            f"| {layer['name']} | {layer['status']} | {layer['objective_value']} | "
            f"{layer['best_bound']} | {layer['gap']} | {layer['seconds']:.3f} |"
        )
    lines += [
        "",
        "## 分类统计",
        "",
        "| 类别 | 保留 | 调整 | 撤销 |",
        "|---|---:|---:|---:|",
    ]
    for category, total in (("A", 20), ("B", 40), ("C", 90)):
        cancelled = metrics.get(f"C_{category}", 0)
        adjusted = metrics.get(f"M_{category}", 0)
        lines.append(f"| {category} | {total - cancelled - adjusted} | {adjusted} | {cancelled} |")
    selected_records = report.get("selected", []) or []
    action_counts = Counter(record.get("action") for record in selected_records)
    lines += [
        "",
        "## 动作与位移",
        "",
        "- actions: "
        + ", ".join(
            f"{name}={action_counts.get(name, 0)}"
            for name in ("keep", "frequency", "time", "gap", "cancel")
        ),
        f"- raw shift sums: `Df_sum={metrics.get('Df_sum')}`, `Dt_sum={metrics.get('Dt_sum')}`, `Dg_sum={metrics.get('Dg_sum')}`",
    ]
    lines += ["", "## 统一复验", ""]
    validation = report.get("validation")
    if validation is None:
        lines += [
            "- 本预算下没有返回可复验的完整排程；不能生成正式 result4。",
            f"- solver status: `{report.get('status', 'UNKNOWN')}`",
        ]
    else:
        lines += [
            f"- plans: `{validation['plan_count']}`",
            f"- occurrences: `{validation['occurrence_count']}`",
            f"- conflicts: `{validation['conflict_count']}`",
            f"- boundary violations: `{validation['boundary_violations']}`",
            f"- fast selected-state conflicts: `{report['selected_pairwise_conflicts']}`",
            f"- canonical ok: `{validation['ok']}`",
        ]
    lines += [
        "",
        "Q4 的 gap 状态仅允许 C 类使用，且与频移、时间平移和撤销互斥；结果不得与 Q2 的排程行混合解释。",
        "若某层状态为 UNKNOWN 或 FEASIBLE，目标值只能作为 incumbent/边界报告，不能写成严格词典序最优。",
        "当前主文件若含 fixed 前缀，只表示在该前缀面上的条件性搜索；Q4 首层 C 的全局最优性仍需单独闭合。",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=ROOT / "D题/附件/附件1.xlsx")
    parser.add_argument("--template", type=Path, default=ROOT / "D题/附件/附件2/result4.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/q4/result4.xlsx")
    parser.add_argument("--summary", type=Path, default=ROOT / "outputs/q4/summary.md")
    parser.add_argument("--report", type=Path, default=ROOT / "outputs/q4/solver_report.json")
    parser.add_argument("--selected", type=Path, default=ROOT / "outputs/q4/selected_T.json")
    parser.add_argument("--candidate-summary", type=Path, default=ROOT / "outputs/q4/candidate_summary.json")
    parser.add_argument("--hint", type=Path, default=ROOT / "outputs/q2/Q2-T-incumbent-v1.json")
    parser.add_argument("--policy", choices=("T", "P"), default="T")
    parser.add_argument("--horizon", type=int, default=643)
    parser.add_argument("--time-limit", type=float, default=180.0)
    parser.add_argument("--workers", type=int, default=16)
    parser.add_argument("--seed", type=int, default=20260912)
    parser.add_argument("--budget-c", type=int, default=None, help="optional cancellation upper bound")
    parser.add_argument("--budget-m", type=int, default=None, help="optional adjustment upper bound")
    parser.add_argument("--fixed-c", type=int, default=None, help="certified cancellation prefix")
    parser.add_argument("--fixed-m", type=int, default=None, help="conditional adjustment prefix")
    args = parser.parse_args()

    plans = load_plans(args.input)
    grouped, flat = generate_all_candidates(plans, horizon=args.horizon)
    edges = generate_conflict_edges(grouped)
    cliques = generate_cell_cliques(flat)
    audit_pairs = edges[:500] + edges[-500:]
    checked = verify_mask_sample(flat, audit_pairs) if audit_pairs else 0
    candidate_summary = {
        "plan_count": len(plans),
        "candidate_state_count": len(flat),
        "candidate_min": min(map(len, grouped)),
        "candidate_max": max(map(len, grouped)),
        "gap_candidate_count": sum(1 for candidate in flat if candidate.action == "gap"),
        "state_conflict_edges": len(edges),
        "cell_clique_count": len(cliques),
        "mask_canonical_sample_checked": checked,
        "horizon": args.horizon,
    }
    _write_json(args.candidate_summary, candidate_summary)

    hints = _load_hint(args.hint, flat) if args.hint.exists() else None
    budgets = {}
    if args.budget_c is not None:
        budgets["C"] = args.budget_c
    if args.budget_m is not None:
        budgets["M"] = args.budget_m
    fixed_initial = {}
    if args.fixed_c is not None:
        fixed_initial["C"] = args.fixed_c
    if args.fixed_m is not None:
        fixed_initial["M"] = args.fixed_m
    fixed_initial = fixed_initial or None
    result = solve_lexicographic(
        grouped,
        flat,
        policy=args.policy,
        mode="full",
        edges=[],
        config=SolverConfig(
            time_limit_seconds=args.time_limit,
            num_search_workers=args.workers,
            random_seed=args.seed,
        ),
        cell_cliques=cliques,
        constraint_form="cells",
        hints=hints,
        budgets=budgets or None,
        fixed_initial=fixed_initial,
    )
    if not result.selected:
        report = {
            "question": "Q4",
            "policy": args.policy,
            "horizon": args.horizon,
            "candidate_summary": candidate_summary,
            "budgets": budgets,
            "fixed_initial": fixed_initial or {},
            "objective_names": list(objective_names(args.policy)),
            "layers": result.layer_dicts(),
            "proven_optimal": result.proven_optimal,
            "status": "NO_FEASIBLE_SELECTION",
        }
        _write_json(args.report, report)
        _write_summary(args.summary, report)
        print(json.dumps(report, ensure_ascii=False))
        return

    selected = list(result.selected)
    validation = validate_schedule([candidate.state for candidate in selected], horizon=args.horizon)
    _write_result4(args.output, args.template, selected)
    metrics = dict(result.metrics)
    report = {
        "question": "Q4",
        "policy": args.policy,
        "input": str(args.input),
        "template": str(args.template),
        "horizon": args.horizon,
        "candidate_summary": candidate_summary,
        "budgets": budgets,
        "fixed_initial": fixed_initial or {},
        "objective_names": list(objective_names(args.policy)),
        "solver_config": {
            "time_limit_seconds_per_layer": args.time_limit,
            "num_search_workers": args.workers,
            "random_seed": args.seed,
            "constraint_form": "cells",
        },
        "metrics": metrics,
        "layers": result.layer_dicts(),
        "selected_pairwise_conflicts": validate_selected_pairwise(selected),
        "selected": [_record(candidate) for candidate in selected],
        "validation": {
            "plan_count": validation.plan_count,
            "occurrence_count": validation.occurrence_count,
            "conflict_count": validation.conflict_count,
            "boundary_violations": validation.boundary_violations,
            "ok": validation.ok,
        },
        "proven_optimal": result.proven_optimal,
    }
    _write_json(args.selected, {"policy": args.policy, "metrics": metrics, "selected": report["selected"], "validation": report["validation"]})
    _write_json(args.report, report)
    _write_summary(args.summary, report)
    print(json.dumps({"policy": args.policy, "metrics": metrics, "proven_optimal": result.proven_optimal, "conflicts": validation.conflict_count, "result": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
